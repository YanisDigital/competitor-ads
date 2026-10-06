#!/usr/bin/env python3
"""Export a snapshot folder (ads.csv + run.json) to an Excel workbook.

    python export_xlsx.py out/<preset>/<date>            # -> <folder>/report.xlsx
    python export_xlsx.py out/<preset>/<date> --out x.xlsx

Needs openpyxl (`pip install openpyxl`; not installed silently). The report
numbers come from collector.js via report.js (node), so this script does not
reimplement any aggregation; it only lays the results out. If node is not
available it falls back to the folder's report.json (some columns stay empty).
If the folder has a diff.json (from compare.js) a "Changes" sheet is added; with
creatives/manifest.json (fetch_creatives.py) a "Creatives" sheet with thumbnails
(embedded only when Pillow is installed) and a summary block of the tags.

Summary counts are Excel formulas (COUNTIF over the ads sheet); Excel computes
them when the file is opened.
"""
import argparse
import csv
import json
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import parse_qs, urlparse

HERE = Path(__file__).parent


class F(str):
    """A formula this script writes on purpose. Any other cell that starts
    with "=" holds third-party text (ad copy, page names, hypotheses) and is
    forced to plain text before saving, so it can never run as a formula."""

try:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required for the Excel export. Install it with: pip install openpyxl")


def load_report(folder: Path):
    try:
        res = subprocess.run(["node", str(HERE / "report.js"), str(folder)], capture_output=True, encoding="utf-8", timeout=120)
        if res.returncode == 0 and res.stdout.strip():
            data = json.loads(res.stdout)
            return data["report"], data["doors"], data["meta"], data.get("query_stats", []), data.get("warnings", []), data.get("eu"), data.get("visuals")
        print("warning: report.js failed, using report.json:", res.stderr.strip(), file=sys.stderr)
    except FileNotFoundError:
        print("warning: node not found, using report.json (door column will be empty)", file=sys.stderr)
    rp = folder / "report.json"
    if not rp.exists():
        sys.exit("No node and no report.json: cannot build the report.")
    meta = json.loads((folder / "run.json").read_text(encoding="utf-8")) if (folder / "run.json").exists() else {}
    return json.loads(rp.read_text(encoding="utf-8")), {}, meta, [], [], None, None


def days_shown(r, run_dt) -> int:
    """Days an ad was shown: until the stop date for a stopped ad, until the collection date otherwise."""
    start = datetime.strptime(r["start"], "%Y-%m-%d")
    if r.get("active") == "false" and r.get("end"):
        return (datetime.strptime(r["end"], "%Y-%m-%d") - start).days
    return (run_dt - start).days


def domain(u: str) -> str:
    try:
        x = urlparse(u)
        if x.hostname and x.hostname.endswith("facebook.com") and parse_qs(x.query).get("u"):
            x = urlparse(parse_qs(x.query)["u"][0])
        h = x.hostname or ""
        for p in ("www.", "l.", "m."):
            if h.startswith(p):
                h = h[len(p):]
        return h
    except Exception:
        return ""


def main() -> None:
    ap = argparse.ArgumentParser(description="Export an ads snapshot folder to Excel.")
    ap.add_argument("folder", help="snapshot folder with ads.csv (e.g. out/ecom-dropship-us/2026-09-30)")
    ap.add_argument("--out", help="output .xlsx path (default: <folder>/report.xlsx)")
    args = ap.parse_args()
    folder = Path(args.folder)
    if not (folder / "ads.csv").exists():
        sys.exit(f"No ads.csv in {folder}")
    out_path = Path(args.out) if args.out else folder / "report.xlsx"

    report, doors, meta, query_stats, warnings, eu, visuals = load_report(folder)
    with open(folder / "ads.csv", encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
    # curation.json (hand-checked list of competitors) was applied by report.js: drop the same advertisers here
    curation = meta.get("curation") or {}
    dropped = set(curation.get("excluded_pages", []))
    if dropped:
        rows = [r for r in rows if r["page"] not in dropped]
    types = curation.get("types", {})
    ts = meta.get("ts") or (folder / "ads.csv").stat().st_mtime
    run_dt = datetime.fromtimestamp(ts, tz=timezone.utc).replace(tzinfo=None)

    FONT = "Arial"
    f_base, f_bold = Font(name=FONT, size=10), Font(name=FONT, size=10, bold=True)
    f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    f_title = Font(name=FONT, size=14, bold=True)
    f_note = Font(name=FONT, size=9, italic=True, color="595959")
    f_link = Font(name=FONT, size=10, color="0563C1", underline="single")
    fill_head = PatternFill("solid", fgColor="1F3864")
    fill_sec = PatternFill("solid", fgColor="D9E1F2")

    def header(ws, row, titles, widths=None):
        for i, t in enumerate(titles, 1):
            c = ws.cell(row=row, column=i, value=t)
            c.font, c.fill = f_head, fill_head
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    own_formulas = set()  # (sheet, row, column) of the formulas written via F()

    def put(ws, row, values):
        for c, v in enumerate(values, 1):
            ws.cell(row=row, column=c, value=v).font = f_base
            if isinstance(v, F):
                own_formulas.add((ws.title, row, c))

    wb = Workbook()

    # ---- Ads ----
    wa = wb.active
    wa.title = "Объявления"
    cols = ["ID", "Ссылка на объявление", "Страница", "Начало показа", "Дней на дату сбора", "Формат", "Вариантов", "CTA",
            "Дверь", "Домен", "Платформы", "Запрос", "Заголовок", "Текст"]
    widths = [18, 34, 30, 13, 12, 12, 10, 16, 20, 26, 26, 22, 40, 90]
    extra = bool(rows) and "cta_type" in rows[0]  # columns added in v0.12.4; older ads.csv files lack them
    if extra:
        cols += ["Тип кнопки", "Подпись ссылки", "Создано ИИ (пометка Meta)"]
        widths += [18, 28, 16]
    status_cols = bool(rows) and "end" in rows[0]  # added with stopped ads (v0.13.0)
    if status_cols:
        cols += ["Статус", "Остановлено"]
        widths += [13, 13]
    media_cols = bool(rows) and "image_url" in rows[0]  # added in v0.13.4: links to the picture and the video
    if media_cols:
        cols += ["Картинка", "Видео"]
        widths += [11, 9]
    media_start = len(cols) - 1  # 1-based column of "Картинка" when media_cols
    header(wa, 1, cols, widths)
    for i, r in enumerate(rows, 2):
        start = datetime.strptime(r["start"], "%Y-%m-%d")
        url = "https://www.facebook.com/ads/library/?id=" + r["id"]
        values = [r["id"], url, r["page"], start, days_shown(r, run_dt), r["fmt"] or "?", int(r["variants"] or 1),
                  r["cta"] or "(нет)", doors.get(r["id"], ""), domain(r["link"]), r["platforms"].replace("|", ", "),
                  r["kws"], r["title"], r["body"]]
        if extra:
            values += [r.get("cta_type", ""), r.get("caption", ""), {"true": "да", "false": "нет"}.get(r.get("ai_made", ""), "")]
        if status_cols:
            values += ["остановлено" if r.get("active") == "false" else "активно", r.get("end", "")]
        shown = {"image": r.get("image_url", "") if media_cols else "", "video": r.get("video_url", "") if media_cols else ""}
        shown = {k: v for k, v in shown.items() if v.lower().startswith(("http://", "https://"))}  # only real web links
        if media_cols:
            values += ["превью" if shown.get("video") and shown.get("image") else "картинка" if shown.get("image") else None,
                       "видео" if shown.get("video") else None]
        put(wa, i, values)
        for offset, key in ((0, "image"), (1, "video")):
            if key in shown:
                cell = wa.cell(row=i, column=media_start + offset)
                cell.hyperlink = shown[key]
                cell.font = f_link
        wa.cell(row=i, column=2).hyperlink = url
        wa.cell(row=i, column=2).font = f_link
        wa.cell(row=i, column=4).number_format = "yyyy-mm-dd"
    last_ad = len(rows) + 1
    if media_cols:
        note = Comment(
            "Ссылки на картинки и видео выдаёт CDN Meta с подписью: они действуют ограниченное время (обычно дни). "
            "Нужное открой или скачай сразу; потом ссылка перестанет работать, а объявление останется доступным по колонке «Ссылка на объявление».",
            "meta-ads-niche-report")
        note.width, note.height = 320, 120
        wa.cell(row=1, column=media_start).comment = note
    wa.freeze_panes = "D2"
    wa.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last_ad}"
    wa.row_dimensions[1].height = 30

    # ---- Advertisers ----
    local_set, platform_set = set(report.get("local_pages", [])), set(report.get("platform_pages", []))
    stores = defaultdict(list)
    for g in report.get("store_groups", []):
        for p in g["pages"]:
            stores[p].append(g["site"])
    clustered = defaultdict(int)
    for cl in report.get("creative_clusters", []):
        for p in cl["pages"]:
            clustered[p] += 1
    by_page = defaultdict(list)
    for r in rows:
        by_page[r["page"]].append(r)
    age_of = lambda a: days_shown(a, run_dt)
    page_rows = []
    for page, ads in by_page.items():
        ages = [age_of(a) for a in ads]
        oldest = max(ads, key=age_of)
        page_rows.append((page, len(ads), max(ages), min(ages), max(int(a["variants"] or 1) for a in ads),
                          ", ".join(f"{k} {v}" if v > 1 else k for k, v in Counter(a["fmt"] or "?" for a in ads).most_common()),
                          ", ".join(sorted({doors.get(a["id"], "") for a in ads} - {""})),
                          ", ".join(sorted({domain(a["link"]) for a in ads if domain(a["link"])})),
                          "да" if page in local_set else "", "да" if page in platform_set else "",
                          "; ".join(stores.get(page, [])), clustered.get(page, 0) or "",
                          "https://www.facebook.com/ads/library/?id=" + oldest["id"]))
    page_rows.sort(key=lambda x: (-x[1], -x[2]))
    wp = wb.create_sheet("Рекламодатели")
    pcols = ["Страница", "Объявлений", "Самое старое, дней", "Самое свежее, дней", "Макс. вариантов", "Форматы", "Двери", "Сайты",
             "Локальный", "Платформа", "Общий сайт с другими страницами", "Копии креативов на др. страницах", "Пример объявления"]
    pwidths = [38, 11, 12, 12, 11, 22, 26, 36, 10, 10, 34, 14, 40]
    likes = {}
    for r in rows:
        if r.get("page_likes", "").isdigit():
            likes[r["page"]] = max(likes.get(r["page"], 0), int(r["page_likes"]))
    if likes:
        pcols.append("Подписчиков страницы")
        pwidths.append(14)
    if types:
        pcols.append("Тип (ручная разметка)")
        pwidths.append(34)
    header(wp, 1, pcols, pwidths)
    for i, pr in enumerate(page_rows, 2):
        vals = list(pr)
        if likes:
            vals.append(likes.get(pr[0]))
        if types:
            vals.append(types.get(pr[0], ""))
        vals[1] = F(f"=COUNTIF('Объявления'!$C$2:$C${last_ad},A{i})")
        put(wp, i, vals)
        wp.cell(row=i, column=13).hyperlink = pr[-1]
        wp.cell(row=i, column=13).font = f_link
    last_page = len(page_rows) + 1
    wp.freeze_panes = "B2"
    wp.auto_filter.ref = f"A1:{get_column_letter(len(pcols))}{last_page}"
    wp.row_dimensions[1].height = 45

    # ---- EU reach and audience (only when eu_details.py was run) ----
    if eu and eu.get("ads"):
        we = wb.create_sheet("ЕС охват и аудитория")
        we.cell(row=1, column=1, value=("Данные Meta только для объявлений, показанных в ЕС. Охват это люди по каждому объявлению; сумма по объявлениям одного рекламодателя "
                                        "считает пересекающихся людей дважды (верхняя граница). Возраст и пол охваченных это фактический состав, таргетинг это настройки рекламодателя.")).font = f_note
        header(we, 3, ["Страница", "Объявл. с данными", "Охват, сумма", "Охват, макс.", "Возраст таргетинга", "Пол", "Страны таргетинга", "Главная возрастная группа", "Плательщик"], [38, 12, 13, 13, 14, 14, 36, 14, 40])
        r0 = 4
        for p in eu["per_page"]:
            put(we, r0, [p["page"], p["ads"], p["reach_sum"], p["reach_max"], f'{p["age_min"]}-{p["age_max"]}' if p["age_min"] is not None else "",
                         ", ".join(p["genders"]), ", ".join(p["countries"]), p["top_age_range"], "; ".join(p["payers"])])
            r0 += 1
        r0 += 1
        we.cell(row=r0, column=1, value="По объявлениям").font = f_bold
        r0 += 1
        for i, t in enumerate(["Страница", "Охват", "Возраст таргетинга", "Пол таргетинга", "Главная группа", "Её доля", "Женщин среди охваченных", "Плательщик", "Ссылка"], 1):
            c = we.cell(row=r0, column=i, value=t)
            c.font, c.fill = f_head, fill_head
        r0 += 1
        for a in eu["per_ad"]:
            put(we, r0, [a["page"], a["reach"], f'{a["age_min"]}-{a["age_max"]}' if a["age_min"] is not None else "", a["gender"], a["top_age_range"],
                         a["top_age_share"], a["female_share"], a["payer"], a["url"]])
            we.cell(row=r0, column=9).hyperlink = a["url"]
            we.cell(row=r0, column=9).font = f_link
            for col in (6, 7):
                we.cell(row=r0, column=col).number_format = "0%"
            r0 += 1
        r0 += 1
        we.cell(row=r0, column=1, value="Возраст охваченных, все объявления").font = f_bold
        r0 += 1
        for k, v in eu["overall"]["age_share"].items():
            put(we, r0, [k, v])
            we.cell(row=r0, column=2).number_format = "0%"
            r0 += 1
        we.freeze_panes = "A4"

    # ---- Stopped ads (only when the run asked for them) ----
    stopped = report.get("stopped") or {}
    if stopped.get("ads"):
        wsd = wb.create_sheet("Остановленные")
        header(wsd, 1, ["Страница", "Дней показа", "Вариантов", "Формат", "Дверь", "Текст (начало)", "Ссылка"], [36, 11, 11, 12, 20, 90, 44])
        for i, t in enumerate(stopped.get("top", []), 2):
            put(wsd, i, [t["page"], t["run_days"], t["variants"], t["fmt"], t["door"], t["text"], t["url"]])
            wsd.cell(row=i, column=7).hyperlink = t["url"]
            wsd.cell(row=i, column=7).font = f_link
        wsd.cell(row=len(stopped.get("top", [])) + 3, column=1, value=(
            f"Остановлено {stopped['ads']} объявл. от {stopped['advertisers']} рекламодателей; медиана показа {stopped['median_run_days']} дн.; "
            f"{stopped['short_lived']} остановлено меньше чем за 2 недели (похоже на тест, который не сработал). Долгий показ до остановки не значит, что объявление работало.")).font = f_note
        wsd.freeze_panes = "A2"

    # ---- Longrun ----
    wl = wb.create_sheet("Долгожители")
    header(wl, 1, ["Страница", "Дней", "Формат", "Вариантов", "Дверь", "Текст (начало)", "Ссылка"], [36, 9, 12, 11, 20, 90, 44])
    for i, l in enumerate(report.get("longrun", []), 2):
        put(wl, i, [l["page"], l["days"], l["fmt"], l["variants"], l["door"], l["text"], l["url"]])
        wl.cell(row=i, column=7).hyperlink = l["url"]
        wl.cell(row=i, column=7).font = f_link
    wl.cell(row=len(report.get("longrun", [])) + 3, column=1,
            value="Ранжирование: по числу вариантов креатива, потом по возрасту; не больше 2 объявлений на страницу; у онлайн-пресетов локальные бизнесы и платформы исключены.").font = f_note
    wl.freeze_panes = "A2"

    # ---- Creatives (only when fetch_creatives.py was run) ----
    if visuals and visuals.get("items"):
        try:
            import PIL  # noqa: F401  (openpyxl embeds pictures through Pillow)
            from openpyxl.drawing.image import Image as XLImage
        except ImportError:
            XLImage = None
        wv = wb.create_sheet("Креативы")
        wv.cell(row=1, column=1, value=(
            f"Что изображено на креативах: скачано {visuals['downloaded']}, размечено {visuals['labeled']} у {visuals['advertisers']} рекламодателей. "
            "Видео оценены только по кадру-превью. Сводка по тегам на листе «Сводка»; сила сигнала считается по рекламодателям, как у хуков.")).font = f_note
        tf = visuals.get("tag_fields", [])
        vcols = ["Миниатюра", "Страница", "Вид", "Дней", "Долгожитель"] + [t["label"] for t in tf] + ["Заметка", "Объявление", "Файл"]
        header(wv, 3, vcols, [18, 30, 12, 8, 11] + [16] * len(tf) + [40, 13, 30])
        root = folder.resolve()
        kind_names = {"image": "картинка", "carousel": "карусель", "video_preview": "кадр видео"}
        for i, g in enumerate(visuals["items"], 4):
            names = g.get("tag_names") or {}
            put(wv, i, [None, g["page"], kind_names.get(g["kind"], g["kind"]), g["days"], "да" if g["long_running"] else ""]
                + [names.get(t["field"], "") for t in tf] + [g.get("notes", "") if g.get("tags") else "не размечено", "объявление", (g.get("files") or [""])[0]])
            for c in range(6, len(vcols) - 1):
                wv.cell(row=i, column=c).alignment = Alignment(wrap_text=True, vertical="top")
            wv.cell(row=i, column=len(vcols) - 1).hyperlink = g["url"]
            wv.cell(row=i, column=len(vcols) - 1).font = f_link
            if XLImage is None:
                continue
            for rel in (g.get("thumbs") or [])[:1] + (g.get("files") or [])[:1]:
                p = (folder / rel).resolve()
                if not p.is_relative_to(root) or not p.is_file() or p.suffix.lower() not in (".jpg", ".jpeg", ".png", ".gif"):
                    continue  # stay inside the snapshot folder
                try:
                    img = XLImage(str(p))
                except Exception:
                    continue
                w = 120
                img.height, img.width = max(1, round(img.height * w / max(1, img.width))), w
                img.anchor = f"A{i}"
                wv.add_image(img)
                wv.row_dimensions[i].height = max(wv.row_dimensions[i].height or 15, img.height * 0.75 + 4)
                break
        wv.freeze_panes = "C4"

    # ---- Networks ----
    wn = wb.create_sheet("Сети страниц")
    wn.cell(row=1, column=1, value="Один и тот же текст объявления на разных страницах").font = f_bold
    header(wn, 2, ["Текст (начало)", "Страниц", "Объявлений", "Самое старое, дней", "Страницы"], [70, 10, 11, 14, 90])
    r0 = 3
    for cl in report.get("creative_clusters", []):
        put(wn, r0, [cl["text"], len(cl["pages"]), cl["ads"], cl["oldest_days"], "; ".join(cl["pages"])])
        r0 += 1
    r0 += 1
    wn.cell(row=r0, column=1, value="Разные страницы, ведущие на один сайт").font = f_bold
    r0 += 1
    for i, t in enumerate(["Сайт", "Страниц", "", "", "Страницы"], 1):
        c = wn.cell(row=r0, column=i, value=t)
        c.font, c.fill = f_head, fill_head
    r0 += 1
    for g in report.get("store_groups", []):
        put(wn, r0, [g["site"], len(g["pages"]), None, None, "; ".join(g["pages"])])
        r0 += 1

    # ---- Queries: which ones find competitors ----
    if query_stats:
        curated = any(q.get("relevant_advertisers") is not None for q in query_stats)
        wq = wb.create_sheet("Запросы")
        qcols = ["Запрос", "Объявлений", "Рекламодателей"] + (["Из них конкуренты (после ручной проверки)"] if curated else []) + ["Лимит Meta"]
        header(wq, 1, qcols, [36, 12, 15] + ([20] if curated else []) + [14])
        for i, q in enumerate(query_stats, 2):
            put(wq, i, [q["query"], q["ads"], q["advertisers"]] + ([q["relevant_advertisers"]] if curated else []) + ["да" if q.get("rate_limited") else ""])
        wq.cell(row=len(query_stats) + 3, column=1, value=(
            "Лимит Meta: подгрузка следующих страниц отклонена, по запросу только первая партия (~30 объявлений). "
            "Запросы с нулём рекламодателей или без конкурентов стоит заменить.")).font = f_note
        wq.freeze_panes = "A2"

    # ---- Changes (only if compare.js was run) ----
    diff_path = folder / "diff.json"
    if diff_path.exists():
        d = json.loads(diff_path.read_text(encoding="utf-8"))
        wc = wb.create_sheet("Изменения")
        wc.column_dimensions["A"].width = 34
        for col, w in zip("BCDEF", (12, 12, 12, 70, 44)):
            wc.column_dimensions[col].width = w
        wc["A1"] = f"Сравнение с предыдущим срезом ({d['interval_days']} дн.)"
        wc["A1"].font = f_bold
        summary = [("Объявлений раньше → сейчас", f"{d['prev_ads']} → {d['curr_ads']}"), ("Пережили", d["survived"]),
                   ("Новых", d["new_ads"]["count"]),
                   ("Пропало из выдачи", d["stopped"]["count"]), ("…из них уверенно остановлены", d["stopped"]["high_confidence"]),
                   ("Молодых тестов (<30 дн.) пропало", f"{d['young_tests']['gone']} из {d['young_tests']['ads']}")]
        for i, (k, v) in enumerate(summary, 3):
            put(wc, i, [k, v])
        r0 = 3 + len(summary) + 1
        wc.cell(row=r0 - 1, column=1, value="Пропавшее объявление уверенно остановлено, только если его запрос перезапущен и вернул <90 объявлений; иначе оно могло выпасть из топа выдачи.").font = f_note

        def table(title, cols_, rows_):
            nonlocal r0
            wc.cell(row=r0, column=1, value=title).font = f_bold
            r0 += 1
            for i, t in enumerate(cols_, 1):
                c = wc.cell(row=r0, column=i, value=t)
                c.font, c.fill = f_head, fill_head
            r0 += 1
            for row_ in rows_:
                put(wc, r0, row_)
                r0 += 1
            r0 += 1

        table("Пропало из выдачи", ["Страница", "Уверенность", "Дней к 1-му срезу", "Вариантов", "Текст", "Ссылка"],
              [[s["page"], s["confidence"], s["age_at_prev"], s["variants"], s["text"], s["url"]] for s in d["stopped"]["top"]])
        table("Масштабируются (больше вариантов креатива)", ["Страница", "Было", "Стало", "", "", "Ссылка"],
              [[s["page"], s["variants_prev"], s["variants_curr"], None, None, s["url"]] for s in d["scaling"]])
        table("Новые объявления (с самыми большими сериями)", ["Страница", "Вариантов", "", "", "Текст", "Ссылка"],
              [[s["page"], s["variants"], None, None, s["text"], s["url"]] for s in d["new_ads"]["top"]])
        table("Новые страницы", ["Страница", "Объявлений"], [[p["page"], p["ads"]] for p in d["pages"]["new"]])
        table("Страницы, которые пропали", ["Страница"], [[p] for p in d["pages"]["gone"]])
        table("Страницы, выросшие на 3+ объявления", ["Страница", "Было", "Стало"], [[p["page"], p["prev"], p["curr"]] for p in d["pages"]["grew"]])

    # ---- Hypotheses (only if Claude saved hypotheses.json next to the snapshot) ----
    hyp_path = folder / "hypotheses.json"
    if hyp_path.exists():
        hyps = json.loads(hyp_path.read_text(encoding="utf-8"))
        wh = wb.create_sheet("Гипотезы")
        hcols = [("name", "Название", 26), ("signal_strength", "Сила сигнала", 22), ("angle", "Угол", 34), ("evidence", "Свидетельство из данных", 44),
                 ("headline", "Заголовок", 30), ("primary_text", "Основной текст", 60), ("cta", "CTA", 14),
                 ("destination", "Куда ведёт", 28), ("format", "Формат креатива", 34), ("test", "Что тестируем", 30),
                 ("metric", "Метрика", 20), ("risk", "Риски по политикам", 34), ("confirm_with_client", "Подтвердить у клиента", 34), ("lint", "Автопроверка", 60)]
        lint_path = folder / "hypotheses_lint.json"
        lint = {}
        if lint_path.exists():
            for r in json.loads(lint_path.read_text(encoding="utf-8")).get("results", []):
                lint[r["name"]] = "\n".join(f'[{f["severity"]}] {f["message"]}' for f in r["findings"] if f["severity"] != "info")
        header(wh, 1, [h[1] for h in hcols], [h[2] for h in hcols])
        for i, h in enumerate(hyps, 2):
            for c, (key, _, _) in enumerate(hcols, 1):
                v = lint.get(h.get("name"), "") if key == "lint" else h.get(key, "")
                cell = wh.cell(row=i, column=c, value="; ".join(map(str, v)) if isinstance(v, list) else v)
                cell.font = f_base
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        wh.freeze_panes = "B2"
        wh.row_dimensions[1].height = 30

    # ---- Test plan (only if plan_tests.js was run) ----
    plan_path = folder / "test_plan.json"
    if plan_path.exists():
        tp = json.loads(plan_path.read_text(encoding="utf-8"))
        wt = wb.create_sheet("План тестов")
        header(wt, 1, ["#", "Гипотеза", "Балл", "Сила сигнала", "Факт клиента", "Трудоёмкость", "Тип переменной", "Блокеры"], [5, 60, 8, 14, 14, 13, 16, 60])
        for i, r in enumerate(tp["ranked"], 2):
            put(wt, i, [r["rank"], r["name"], r["score"], r["evidence_strength"], r["fit"], r["effort"], r["variable_type"], "; ".join(r["blockers"])])
        r0 = len(tp["ranked"]) + 4
        for i, t in enumerate(["Раунд", "Тесты (параллельно)", "Бюджет"], 1):
            c = wt.cell(row=r0, column=i, value=t)
            c.font, c.fill = f_head, fill_head
        for rd in tp["plan"]["rounds"]:
            r0 += 1
            put(wt, r0, [rd["round"], "; ".join(f'{t["name"]} ({t["variable_type"]})' for t in rd["tests"]), rd["budget"]])
        a = tp["plan"]["assumptions"]
        r0 += 2
        wt.cell(row=r0, column=1, value=(
            f"Допущения: {a['variants_per_test']} варианта на тест, {a['events_per_variant']} событий оптимизации на вариант (ориентир Meta для фазы обучения, проверь), "
            f"не меньше {a['min_days_per_round']} дней на раунд, целевой CPA {a['target_cpa'] if a['target_cpa'] is not None else 'неизвестен (добавь target_cpa в client.json)'}. "
            "Балл = сила сигнала × готовность / трудоёмкость: он упорядочивает работу, а не предсказывает результат.")).font = f_note
        wt.freeze_panes = "A2"

    # ---- Summary ----
    ws = wb.create_sheet("Сводка", 0)
    for col, w in zip("ABCD", (44, 14, 12, 60)):
        ws.column_dimensions[col].width = w
    title = meta.get("preset_title") or meta.get("preset") or "Срез Ads Library"
    ws["A1"], ws["A1"].font = f"Ads Library: {title}", f_title
    queries = meta.get("queries") or sorted({k for r in rows for k in r["kws"].split("; ") if k})
    ws["A2"] = f"Дата сбора: {run_dt:%Y-%m-%d}. Страна: {meta.get('country', '?')}. Запросов: {len(queries)}."
    ws["A3"] = "Снимок на дату сбора: библиотека отдаёт до ~120 объявлений на запрос (верх выдачи по охвату), это не весь рынок."
    ws["A2"].font = ws["A3"].font = f_note
    ws["A4"] = "Запросы: " + " | ".join(queries)
    ws["A4"].font = f_note
    if curation:
        ws["A5"] = (f"Проверено вручную (curation.json): убрано {curation.get('excluded_ads', 0)} объявл. от "
                    f"{len(curation.get('excluded_pages', []))} рекламодателей как не конкурентов, осталось {curation.get('kept_pages', 0)}."
                    + (" " + curation["notes"] if curation.get("notes") else ""))
        ws["A5"].font = f_note
    row = 6
    if warnings:
        ws.cell(row=row, column=1, value="Предупреждения").font = f_bold
        row += 1
        for w in warnings:
            c = ws.cell(row=row, column=1, value="• " + w["message"])
            c.font = f_note
            row += 1
        row += 1

    def section(t):
        nonlocal row
        for c in range(1, 4):
            ws.cell(row=row, column=c).fill = fill_sec
        ws.cell(row=row, column=1, value=t).font = f_bold
        ws.cell(row=row, column=2, value="Кол-во").font = f_bold
        ws.cell(row=row, column=3, value="Доля").font = f_bold
        row += 1

    rng = lambda col: f"'Объявления'!${col}$2:${col}${last_ad}"
    section("Общее")
    put(ws, row, ["Объявлений", F(f"=COUNTA({rng('A')})")])
    total = f"$B${row}"
    row += 1
    put(ws, row, ["Рекламодателей (страниц)", F(f"=COUNTA('Рекламодатели'!$A$2:$A${last_page})")])
    row += 1
    put(ws, row, ["Из них с одним объявлением", F(f"=COUNTIF('Рекламодатели'!$B$2:$B${last_page},1)"), F(f"=B{row}/B{row-1}")])
    ws.cell(row=row, column=3).number_format = "0%"
    row += 2

    def block(t, col, items):
        nonlocal row
        section(t)
        for label in items:
            put(ws, row, [label, F(f"=COUNTIF({rng(col)},A{row})"), F(f"=B{row}/{total}")])
            ws.cell(row=row, column=3).number_format = "0%"
            row += 1
        row += 1

    block("Формат", "F", [k for k, _ in Counter(r["fmt"] or "?" for r in rows).most_common()])
    block("Дверь (куда ведёт)", "I", [k for k, _ in Counter(doors.get(r["id"], "") for r in rows).most_common() if k])
    block("CTA (топ-8)", "H", [k for k, _ in Counter(r["cta"] or "(нет)" for r in rows).most_common(8)])
    section("Возраст объявления")
    for label, lo, hi in [("меньше 7 дней", 0, 6), ("7–30 дней", 7, 29), ("30–90 дней", 30, 89), ("90–365 дней", 90, 364), ("больше 365 дней", 365, 100000)]:
        put(ws, row, [label, F(f'=COUNTIFS({rng("E")},">={lo}",{rng("E")},"<={hi}")'), F(f"=B{row}/{total}")])
        ws.cell(row=row, column=3).number_format = "0%"
        row += 1
    row += 1
    section("Хуки (по regex, значения из отчёта)")
    ws.cell(row=row - 1, column=4, value="Если слова хука входят в запросы, цифра круговая.").font = f_note
    for k, v in sorted(report.get("hook_freq", {}).items(), key=lambda kv: -kv[1]):
        put(ws, row, [k, v, F(f"=B{row}/{total}")])
        ws.cell(row=row, column=3).number_format = "0%"
        row += 1
    row += 1
    if visuals and visuals.get("fields"):
        section(f"Что на креативах (размечено {visuals['labeled']} у {visuals['advertisers']} рекламодателей)")
        ws.cell(row=row - 1, column=4, value="Доля среди размеченных; рекламодатели и сила сигнала: weak = привычка одного-двух, не тренд.").font = f_note
        for f in visuals["fields"]:
            ws.cell(row=row, column=1, value=f["label"]).font = f_bold
            row += 1
            for x in f["values"]:
                put(ws, row, ["   " + x["name"], x["creatives"], x["share"], f"{x['advertisers']} рекл." + (f" · {x['strength']}" if x.get("strength") else "") + (f" · долгожителей {x['long_running']}" if x["long_running"] else "")])
                ws.cell(row=row, column=3).number_format = "0%"
                row += 1
        row += 1
    pr = report.get("prices")
    if pr:
        section(f"Цены ({pr.get('currency', '')}, только объявления с ценой в тексте)")
        for label, key in [("Объявлений с ценой", "ads_with_price"), ("Минимум", "min"), ("Медиана", "median"), ("Максимум", "max"),
                           ("Упоминаний «% off»", "pct_off_mentions"), ("Медианная скидка, %", "median_pct_off")]:
            put(ws, row, [label, pr.get(key)])
            row += 1
    ws.sheet_view.showGridLines = False
    for sheet in wb:
        for cells in sheet.iter_rows():
            for c in cells:
                if c.data_type == "f" and (sheet.title, c.row, c.column) not in own_formulas:
                    c.data_type = "s"  # third-party text that starts with "="
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    print(f"Saved: {out_path} ({len(rows)} ads, {len(page_rows)} advertisers)")


if __name__ == "__main__":
    main()
